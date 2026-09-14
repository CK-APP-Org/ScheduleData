"""Parse the cafeteria weekly menu (週菜單 sheet) from the 熱食部 xlsx file."""
from datetime import datetime, timedelta

import openpyxl

# First column (1-based) of each weekday block: 項次 | 餐別 | 價格 | ...
DAY_START_COLUMNS = {0: 1, 1: 7, 2: 13, 3: 19, 4: 25}  # A, G, M, S, Y
HEADER_SCAN_ROWS = 20


def _find_monday(sheet) -> datetime:
    for row in range(1, HEADER_SCAN_ROWS + 1):
        value = sheet.cell(row=row, column=1).value
        if isinstance(value, str) and value.count("/") == 2:
            try:
                value = datetime(*map(int, value.strip().split("/")))
            except ValueError:
                continue
        if isinstance(value, datetime):
            if value.weekday() != 0:
                raise ValueError(f"Date in column A ({value:%Y-%m-%d}) is not a Monday")
            return value.replace(hour=0, minute=0, second=0, microsecond=0)
    raise ValueError("Could not find Monday date in column A")


def _parse_price(value):
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str) and value.strip():
        text = value.strip()
        return int(text) if text.isdigit() else text
    return None


def scrape_menu(file_path, sheet_name="週菜單"):
    """
    Returns a list of dicts: date, day, dish_number, dish_name, price.
    Rows whose 項次 cell is not a number (e.g. ★, side dishes, notes) are skipped.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found (sheets: {wb.sheetnames})")
    sheet = wb[sheet_name]
    monday = _find_monday(sheet)

    menu_data = []
    for day_offset, col in DAY_START_COLUMNS.items():
        current_date = monday + timedelta(days=day_offset)

        item_row = next(
            (r for r in range(1, sheet.max_row + 1) if sheet.cell(row=r, column=col).value == "項次"),
            None,
        )
        if item_row is None:
            continue

        for row in range(item_row + 1, sheet.max_row + 1):
            number = sheet.cell(row=row, column=col).value
            if isinstance(number, bool) or not isinstance(number, (int, float)):
                continue
            dish = sheet.cell(row=row, column=col + 1).value
            if dish is None or not str(dish).strip():
                break
            menu_data.append({
                "date": current_date,
                "day": current_date.strftime("%A"),
                "dish_number": int(number),
                "dish_name": str(dish).strip(),
                "price": _parse_price(sheet.cell(row=row, column=col + 2).value),
            })

    return menu_data


if __name__ == "__main__":
    import sys

    for item in scrape_menu(sys.argv[1] if len(sys.argv) > 1 else "menu.xlsx"):
        print(f"{item['date']:%Y-%m-%d} {item['dish_number']}. {item['dish_name']} - {item['price']}")
